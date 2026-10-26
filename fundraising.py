#!/usr/bin/env python3
"""Build docs/fundraising.html, the House battlefield fundraising dashboard, from the FEC Master Sheet.

    pip install openpyxl
    python3 fundraising.py "FEC Master Sheet- 2026 Q3.xlsx"

Reads the candidate rows on the regional research tabs, attaches DCCC/NRCC programs (lists below),
takes incumbency and outside-spending totals from docs/data.json (the IE tracker), and writes
docs/fundraising.html. On the live site the page re-reads docs/data.json, so outside spending stays
current as the tracker updates. Re-run this script when there's a new quarter's sheet.

The spreadsheet itself is never committed: only the figures the page shows end up in docs/.
"""
import argparse, json, re, sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SHEETS = ["Northeast", "Mid-Atlantic", "Midwest", "Southern", "Mountain West", "Western", "Frontier", "Horizon"]

# DCCC programs: dccc.org/2026-frontline and dccc.org/2026-red-to-blue (Oct 16, 2026).
# Any race on neither list is "Districts in Play".
FRONTLINE = ("AL-02 CA-13 CA-45 CT-05 FL-09 FL-14 FL-25 IN-01 MI-08 NC-01 NJ-09 NM-02 NV-01 NV-03 NV-04 NY-03 NY-04 "
             "NY-19 OH-01 OH-09 OH-13 OR-05 TX-28 TX-34 VA-07 WA-03").split()
RED_TO_BLUE = ("AK-AL AR-02 AZ-01 AZ-02 AZ-06 CA-06 CA-22 CA-48 CO-03 CO-05 CO-08 FL-07 FL-13 FL-22 FL-27 IA-01 IA-02 "
               "IA-03 KY-06 ME-02 MI-04 MI-07 MI-10 MN-01 MT-01 NC-11 NE-02 NJ-07 NY-17 NY-21 NV-02 OH-07 PA-01 PA-07 "
               "PA-08 PA-10 SC-01 TX-15 TX-35 TN-05 TN-09 VA-01 VA-02 VA-05 WI-01 WI-03").split()
# NRCC programs: the NRCC Patriots and MAGA Majority sheets (Oct 2026).
PATRIOTS = "AK-AL AZ-06 CA-01 CA-22 CO-08 IA-01 IA-03 MI-04 MI-07 NJ-07 NY-17 PA-01 PA-07 PA-08 PA-10 VA-02 WI-03".split()
MAGA_MAJORITY = ("AZ-01 CA-13 CA-48 FL-09 FL-14 FL-22 FL-25 IA-02 IN-01 ME-02 MI-10 MT-01 NC-01 NC-11 NE-02 NH-01 "
                 "NM-02 NV-01 NV-03 NY-03 NY-04 OH-01 OH-09 TX-28 TX-34 TX-35 WA-03").split()

# Which side a candidate counts on, matching config.toml's party overrides.
SIDE_OVERRIDE = {"Kevin Kiley": ("R", "Filed with the FEC as an independent; counted with Republicans."),
                 "Bill Hill": ("D", "Independent; counted on the Democratic side.")}
FORCE_OTHER = {"Eric Hafner"}
# Incumbents the IE tracker doesn't flag (races with no IEs yet, or a name mismatch).
EXTRA_INC = {"Jahana Hayes", "Nick Begich", "Ann Wagner", "Derek Schmidt", "Brad Knott", "Stephanie Bice",
             "Nick LaLota", "Laura Gillen", "Michael Baumgartner", "John McGuire"}


def num(v):
    if v is None: return None
    if isinstance(v, str):
        try: return float(v.replace("$", "").replace(",", ""))
        except ValueError: return None
    return float(v)


def norm(h): return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def read_sheet(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out, seen = [], set()
    for sh in SHEETS:
        if sh not in wb.sheetnames:
            print(f"  ! no '{sh}' tab", file=sys.stderr); continue
        ws = wb[sh]
        hdr = [norm(c.value) for c in ws[1]]

        def col(*pats):
            for i, h in enumerate(hdr):
                if all(p in h for p in pats): return i
            return None

        first = lambda *opts: next((c for c in (col(*o) if isinstance(o, tuple) else col(o) for o in opts) if c is not None), None)
        ix = dict(dist=col("district"), party=col("party"), status=col("status"),
                  pp_rec=first("pre-primary rece", "pre-primary rec"), pp_disb=col("pre-primary disb"),
                  pp_unit=col("pre-primary unitem"),
                  oct_rec=first("october quarterly receipts", "q3 october receipts"),
                  oct_disb=first("october quarterly disb", "q3 october disb"),
                  ctd_rec=col("cycle to date rec"), ctd_disb=first("cycle to date disb", "cycle to date disbursments"),
                  coh=col("cash on hand"), debt=col("debt"), loans=col("loans by candidate"),
                  oct_unit=first("october quarterly unitem", "october unitem"))
        for row in ws.iter_rows(min_row=2):
            name = row[0].value
            d = row[ix["dist"]].value if ix["dist"] is not None else None
            if not name or not d or not re.match(r"^[A-Z]{2}-(\d\d|AL)$", str(d).strip()): continue
            g = lambda k: num(row[ix[k]].value) if ix.get(k) is not None else None
            link = None
            for c in row:
                h = hdr[c.column - 1] if c.column - 1 < len(hdr) else ""
                if "pre-primary" in h: continue
                t = c.hyperlink.target if c.hyperlink and c.hyperlink.target else c.value if isinstance(c.value, str) else ""
                if t and "docquery.fec.gov" in t: link = t; break
            pp_rec, oct_rec = g("pp_rec") or 0, g("oct_rec")
            status = str(row[ix["status"]].value or "").strip() if ix["status"] is not None else ""
            filed = oct_rec is not None and status.lower() != "not filed"
            c = dict(dist=str(d).strip(), party=str(row[ix["party"]].value or "").strip().upper(),
                     name=re.sub(r"\s+", " ", str(name)).strip(), filed=filed,
                     q3=(pp_rec + oct_rec) if filed else None,
                     q3_disb=((g("pp_disb") or 0) + (g("oct_disb") or 0)) if filed else None,
                     pp=pp_rec if filed else None, oct=oct_rec if filed else None,
                     ctd=g("ctd_rec") if filed else None, ctd_disb=g("ctd_disb") if filed else None,
                     coh=g("coh") if filed else None, debt=g("debt"), loans=g("loans"),
                     unit=((g("pp_unit") or 0) + (g("oct_unit") or 0)) if filed else None, link=link)
            key = (c["dist"], c["name"].lower())
            if key in seen: continue
            seen.add(key); out.append(c)
    return out


def build(xlsx, ie_path, tpl_path, out_path):
    cands = read_sheet(xlsx)
    ie = json.loads(Path(ie_path).read_text())
    iebyd = {r["district"]: r for r in ie["races"]}
    last = lambda n: re.sub(r"[^a-z]", "", n.lower().replace("jr.", "").split()[-1]) if n else ""
    r0 = lambda x: None if x is None else round(x)
    by_dist = {}
    for c in cands: by_dist.setdefault(c["dist"], []).append(c)
    races = []
    for dist in sorted(by_dist):
        e = iebyd.get(dist[:2] + "00" if dist.endswith("-AL") else dist.replace("-", ""))
        R = D = None; others = []
        for c in by_dist[dist]:
            inc = c["name"] in EXTRA_INC
            if e:
                for s in (e.get("rep"), e.get("dem")):
                    if s and s.get("incumbent") and last(c["name"]) in re.sub(r"[^a-z ]", "", s["name"].lower()).split():
                        inc = True
            o = dict(n=c["name"], p=c["party"], inc=inc, f=c["filed"], q3=r0(c["q3"]), pp=r0(c["pp"]), oq=r0(c["oct"]),
                     sp=r0(c["q3_disb"]), ctd=r0(c["ctd"]), ctds=r0(c["ctd_disb"]), coh=r0(c["coh"]),
                     debt=r0(c["debt"]) or 0, loans=r0(c["loans"]) or 0, un=r0(c["unit"]), link=c["link"])
            side, note = SIDE_OVERRIDE.get(c["name"], (c["party"], None))
            if note: o["note"] = note
            if c["name"] in FORCE_OTHER: side = "O"
            if side == "R" and R is None: R = o
            elif side == "D" and D is None: D = o
            else: others.append(o)
        races.append(dict(d=dist,
                          pr="FL" if dist in FRONTLINE else "R2B" if dist in RED_TO_BLUE else "DIP",
                          rp="PAT" if dist in PATRIOTS else "MM" if dist in MAGA_MAJORITY else "NONE",
                          ie=dict(r=round(e["r_total"]), d=round(e["d_total"])) if e else None,
                          R=R, D=D, o=others))
    data = dict(as_of="Sep 30, 2026", ie_as_of=ie.get("as_of"), races=races)

    tpl = Path(tpl_path).read_text()
    body = tpl.replace("/*__DATA__*/null", json.dumps(data, separators=(",", ":")))
    body = body.replace("<!--NAV_EXTRA-->", '<a href="./">Outside spending</a><a href="late.html">48-hour notices</a>')
    head_end = body.index("</style>") + len("</style>")
    html = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
            + body[:head_end] + "\n</head>\n<body>\n" + body[head_end:] + "\n</body>\n</html>\n")
    Path(out_path).write_text(html)

    print(f"{len(races)} races, {len(cands)} candidates -> {out_path}")
    print("  DCCC:", dict(Counter(r["pr"] for r in races)), " NRCC:", dict(Counter(r["rp"] for r in races)))
    listed = set(FRONTLINE + RED_TO_BLUE + PATRIOTS + MAGA_MAJORITY)
    missing = sorted(listed - set(by_dist))
    if missing: print("  program races not in the sheet:", ", ".join(missing))
    one_sided = [r["d"] for r in races if not (r["R"] and r["D"])]
    if one_sided: print("  races missing a D or R:", ", ".join(one_sided))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="FEC Master Sheet .xlsx")
    ap.add_argument("--ie", default=str(ROOT / "docs" / "data.json"))
    ap.add_argument("--template", default=str(ROOT / "fundraising_template.html"))
    ap.add_argument("--out", default=str(ROOT / "docs" / "fundraising.html"))
    a = ap.parse_args()
    build(a.xlsx, a.ie, a.template, a.out)
